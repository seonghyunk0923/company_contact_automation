"""발굴 파이프라인 스크립트 엔드투엔드 테스트 (가상 데이터).

  pip install -r requirements.txt pytest
  pytest -q
"""
import csv
import json
import subprocess
import sys
from pathlib import Path

import pytest
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parent.parent
SCRIPTS = ROOT / "scripts"


def run(script, *args):
    r = subprocess.run([sys.executable, str(SCRIPTS / script), *map(str, args)],
                       capture_output=True, text=True, cwd=SCRIPTS)
    assert r.returncode == 0, r.stdout + r.stderr
    return r.stdout


def company(name, region="경기", district="예시시", industry="기계·장비", cond="①",
            phone="031-000-0000", email="확인불가", home="https://www.example.com", **kw):
    return {"기업명": name, "시도": region, "시군구": district, "업종": industry, "임직원수": "20명",
            "전화": phone, "이메일": email, "홈페이지": home, "주요사업": "예시 제품", "최근활동": "",
            "통과조건": cond, "중점분야": "아니오", "재검토": "아니오", "예비후보": "아니오",
            "출처사이트": "예시 디렉토리", "출처URL": ["https://example.com/detail/1"], "담당자": "",
            "제안안": "공정 불량 데이터 정리로 품질 개선", "비고": "", **kw}


@pytest.fixture()
def workspace(tmp_path):
    cfg = json.loads((ROOT / "config/discovery.example.json").read_text(encoding="utf-8"))
    cfg["run"]["max_per_run"] = 4
    cfg["caps"].update(per_district=2, per_source=10, per_region=10)
    cfg_path = tmp_path / "discovery_config.json"
    cfg_path.write_text(json.dumps(cfg, ensure_ascii=False), encoding="utf-8")

    rec = tmp_path / "records"
    run("init_records.py", "--config", cfg_path, "--out", rec, "--start-number", 100)
    with open(rec / "기존기업명단.csv", "a", encoding="utf-8", newline="") as f:
        csv.writer(f).writerow([100, "(주)기존회사", "02-111-1111", "old@example.com", "원본시트"])
    progress = json.loads((rec / "탐색진행상황.json").read_text(encoding="utf-8"))
    progress["탈락기업"] = [
        {"기업명": "재검토회사", "사유": "홈페이지 없음 [재검토가능]", "일자": "2026-01-01"},
        {"기업명": "큰회사", "사유": "종업원 80명", "일자": "2026-01-01"},
    ]
    progress["예비후보(통과 가능·상한으로 보류)"] = [{"기업명": "예비회사", "사유": "상한"}]
    (rec / "탐색진행상황.json").write_text(json.dumps(progress, ensure_ascii=False), encoding="utf-8")

    r = tmp_path / "run"
    run("prep_exclude.py", "--config", cfg_path, "--progress", rec / "탐색진행상황.json",
        "--existing", rec / "기존기업명단.csv", "--cumulative", rec / "기업컨택_누적목록.csv", "--out", r)

    results = {
        "A": {"agent": "A", "calls_used": 40, "next_position": "region=41000 pageIndex=3부터",
              "passed": [
                  company("기존회사"),                                       # 제외명단과 중복
                  company("가나정밀", email="a@example.com", phone="031-000-0001"),  # 1순위
                  company("다라전자", district="예시시", phone="031-000-0002"),
                  company("마바공업", district="예시시", phone="031-000-0003"),        # 시군구 상한 초과 후보
                  company("사아식품", region="인천", district="예시구", cond="②",
                          home="공식홈페이지 없음", phone="032-000-0004", 최근활동="2026.01 채용공고"),
              ],
              "rejected": [{"기업명": "큰회사", "사유": "종업원 85명", "출처": "x"}],
              "new_blocked": ["blocked.example(robots)"], "new_sources": []},
        "E": {"agent": "E", "calls_used": 20, "next_position": "사람인 키워드 2쪽부터",
              "passed": [
                  company("재검토회사", region="서울", district="예시구", phone="02-000-0005", 재검토="예"),
                  company("가나 정밀(주)", phone="031-000-0001"),                 # A와 중복(전화)
              ],
              "rejected": [{"기업명": "예비회사", "사유": "재검토: 폐업", "출처": "x"}],
              "new_blocked": [], "new_sources": [{"이름": "새 조합 회원사", "URL": "https://example.org"}]},
    }
    for k, v in results.items():
        (r / f"result_{k}.json").write_text(json.dumps(v, ensure_ascii=False), encoding="utf-8")
    return tmp_path, cfg_path, rec, r


def test_full_pipeline(workspace):
    tmp, cfg, rec, r = workspace
    assert json.loads((r / "recheck.json").read_text(encoding="utf-8"))[0]["기업명"] == "재검토회사"

    run("select_candidates.py", "--config", cfg, "--run", r)
    sel = json.loads((r / "selected.json").read_text(encoding="utf-8"))
    names = [c["기업명"] for c in sel]
    dups = {d["기업명"] for d in json.loads((r / "duplicates.json").read_text(encoding="utf-8"))}
    overflow = json.loads((r / "overflow.json").read_text(encoding="utf-8"))

    assert names[0] == "가나정밀"                      # 이메일·홈페이지 모두 → 1순위
    assert {"기존회사", "가나 정밀(주)"} <= dups        # 제외명단·서브에이전트 간 중복
    assert len(sel) == 4                              # max_per_run
    district = [c for c in sel if (c["시도"], c["시군구"]) == ("경기", "예시시")]
    assert len(district) <= 2                         # per_district 상한
    assert names.index("사아식품") == len(names) - 1   # 조건②는 마지막 순위
    assert len(overflow) == 1

    # 원문 재확인에서 하나 떨어뜨리고 다시 고르면 예비 후보가 채워짐
    (r / "dropped.json").write_text(json.dumps([{"기업명": "재검토회사", "사유": "주소 불일치"}],
                                               ensure_ascii=False), encoding="utf-8")
    run("select_candidates.py", "--config", cfg, "--run", r)
    sel = json.loads((r / "selected.json").read_text(encoding="utf-8"))
    assert "재검토회사" not in [c["기업명"] for c in sel]

    out = tmp / "outputs"
    run("build_xlsx.py", "--config", cfg, "--run", r, "--progress", rec / "탐색진행상황.json",
        "--existing", rec / "기존기업명단.csv", "--out-dir", out)
    xlsx = next(out.glob("*.xlsx"))
    assert "(101-" in xlsx.name                         # 기존 마지막번호 100 다음
    wb = load_workbook(xlsx)
    assert wb.sheetnames == ["추가분", "검증정보", "탈락", "요약"]
    first = [c.value for c in wb["추가분"][2]]
    assert first[0] == 101 and first[3] == "가나정밀"
    assert first[8] == "031-000-0001 / a@example.com"

    new = tmp / "records_new"
    run("update_records.py", "--config", cfg, "--run", r, "--progress", rec / "탐색진행상황.json",
        "--existing", rec / "기존기업명단.csv", "--cumulative", rec / "기업컨택_누적목록.csv", "--out", new)
    p = json.loads((new / "탐색진행상황.json").read_text(encoding="utf-8"))
    assert p["마지막번호"] == 100 + len(sel)
    assert p["출처별_진행"]["work.go.kr 강소기업"]["다음_위치"].startswith("region=41000")
    assert "새 조합 회원사" in p["출처별_진행"]
    assert "blocked.example(robots)" in p["조회불가_사이트"]
    rj = {x["기업명"]: x for x in p["탈락기업"]}
    assert rj["큰회사"]["사유"] == "종업원 85명"          # 같은 이름은 갱신
    assert rj["재검토회사"]["사유"].startswith("원문 재확인 불일치")
    reserve = [x["기업명"] for x in p["예비후보(통과 가능·상한으로 보류)"]]
    assert "예비회사" not in reserve                     # 다시 탈락 → 예비에서 삭제
    with open(new / "기존기업명단.csv", encoding="utf-8-sig") as f:
        rows = list(csv.DictReader(f))
    assert len(rows) == 1 + len(sel) and rows[-1]["번호"] == str(100 + len(sel))
    with open(new / "기업컨택_누적목록.csv", encoding="utf-8-sig") as f:
        cum = list(csv.DictReader(f))
    assert cum[0]["임직원수"] == "20"


def test_check_dup_short_names(tmp_path):
    ex = tmp_path / "exclude_names.txt"
    ex.write_text("# [NAMES]\n(주)일\t일\n(주)케이\t케이\n가나정밀\t가나정밀\n# [PHONES]\n# [EMAILS]\n",
                  encoding="utf-8")
    out = run("check_dup.py", "--exclude", ex, "주식회사 동일공사", "케이아이산업개발", "가나정밀(안산)")
    assert "동일공사 => OK" in out
    assert "케이아이산업개발 => OK" in out
    assert "가나정밀(안산) => 중복/제외" in out
