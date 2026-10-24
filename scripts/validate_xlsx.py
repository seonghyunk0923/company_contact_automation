"""만들어진 추가분 엑셀을 다시 열어 형식을 검증한다. (build_xlsx.py가 저장 직후 자동으로 부름)

확인 항목
  - 시트 4개(추가분·검증정보·탈락·요약) 존재
  - '추가분' 머리글 16열이 컨택 시트와 정확히 일치
  - 팀원이 채울 열(팀원명·컨택일자 등 9개)이 모두 빈칸
  - 번호가 1씩 연속, 기업명 중복 없음, 최대 행 수 이하
  - 담당자 연락처가 텍스트인지(앞자리 0 손실 방지)

사용법: python3 scripts/validate_xlsx.py "기업컨택_리스트_추가분_YYYYMMDD_HHmm(시작-끝).xlsx" [--max-rows 50]
문제가 있으면 내용을 출력하고 종료 코드 1로 끝난다.
"""
import argparse
import sys

from openpyxl import load_workbook

from schema import CONTACT_HEADERS, MUST_BE_BLANK, SHEET_NAMES


def validate(path: str, max_rows: int = 50) -> list:
    errors = []
    wb = load_workbook(path)
    missing = [s for s in SHEET_NAMES if s not in wb.sheetnames]
    if missing:
        errors.append(f"시트 없음: {missing}")
    if SHEET_NAMES[0] not in wb.sheetnames:
        return errors

    ws = wb[SHEET_NAMES[0]]
    header = [c.value for c in ws[1]][: len(CONTACT_HEADERS)]
    if header != CONTACT_HEADERS:
        errors.append(f"머리글 불일치: {header}")
        return errors
    col = {h: i for i, h in enumerate(CONTACT_HEADERS)}

    rows = [r for r in ws.iter_rows(min_row=2, values_only=True) if any(v not in (None, "") for v in r)]
    if not rows:
        errors.append("데이터 행 없음")
        return errors
    if len(rows) > max_rows:
        errors.append(f"행 수 {len(rows)} > {max_rows}")

    numbers = [r[col["번호"]] for r in rows]
    if any(not isinstance(n, int) for n in numbers):
        errors.append("번호에 정수가 아닌 값이 있음")
    elif numbers != list(range(numbers[0], numbers[0] + len(numbers))):
        errors.append(f"번호가 연속되지 않음: {numbers[0]}~{numbers[-1]}")

    names = [r[col["기업명"]] for r in rows]
    dups = {n for n in names if names.count(n) > 1}
    if dups:
        errors.append(f"기업명 중복: {sorted(dups)}")

    for i, r in enumerate(rows, start=2):
        filled = [h for h in MUST_BE_BLANK if r[col[h]] not in (None, "")]
        if filled:
            errors.append(f"{i}행: 빈칸이어야 할 열에 값 있음 {filled}")
        contact = r[col["담당자 연락처(전화번호+이메일)"]]
        if contact is not None and not isinstance(contact, str):
            errors.append(f"{i}행: 담당자 연락처가 텍스트가 아님({type(contact).__name__})")
    return errors


if __name__ == "__main__":
    ap = argparse.ArgumentParser()
    ap.add_argument("path")
    ap.add_argument("--max-rows", type=int, default=50)
    a = ap.parse_args()
    problems = validate(a.path, a.max_rows)
    if problems:
        print("검증 실패")
        for p in problems:
            print(" -", p)
        sys.exit(1)
    print("검증 통과")
