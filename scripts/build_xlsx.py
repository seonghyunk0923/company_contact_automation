"""3단계: 선정 결과로 추가분 엑셀(4개 시트)을 만들고 바로 검증한다.

  python3 scripts/build_xlsx.py --config discovery_config.json --run run/ \
      --progress records/탐색진행상황.json --existing records/기존기업명단.csv --out-dir outputs/

시트
  추가분    컨택 시트와 같은 16열. 팀원·컨택 관련 9개 열은 빈칸
  검증정보  기업별 시도·시군구·업종·통과조건·출처 URL·확인 필요 사항
  탈락      서브에이전트 탈락 + 원문 재확인 탈락 + 중복 제거
  요약      summary.json 내용

번호는 진행기록 '마지막번호'와 기존기업명단 최대 번호 중 큰 값 + 1부터 매긴다.
번호를 붙인 선정 결과는 run/numbered.json 으로 저장해 4단계(update_records.py)가 쓴다.
"""
import argparse
import csv
import sys
from datetime import datetime, timedelta, timezone
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font

from schema import (
    CONTACT_HEADERS, K_LAST_NO, NO_HOMEPAGE, REJECT_HEADERS, SHEET_NAMES, UNKNOWN, VERIFY_HEADERS,
    has_value, load_config, load_json, save_json,
)
from select_candidates import is_focus, load_candidates, tier
from validate_xlsx import validate

KST = timezone(timedelta(hours=9))


def next_number(progress_path, existing_path):
    last = 0
    if progress_path and Path(progress_path).exists():
        last = int(load_json(progress_path).get(K_LAST_NO) or 0)
    if existing_path and Path(existing_path).exists():
        with open(existing_path, encoding="utf-8-sig", newline="") as f:
            for row in csv.DictReader(f):
                try:
                    last = max(last, int(str(row.get("번호", "")).strip()))
                except ValueError:
                    pass
    return last + 1


def val(v, empty=UNKNOWN):
    return str(v).strip() if has_value(v) else empty


def intro(c):
    """'기업소개 및 제안주제' 칸"""
    urls = []
    for u in c.get("출처URL") or []:
        if u and u not in urls:
            urls.append(u)
    parts = [
        f"소재지 : {c.get('시도', '')} {c.get('시군구', '')}".strip(),
        f"임직원 수: {val(c.get('임직원수'))}",
        f"업종: {c.get('업종', '')}",
        f"주요사업: {val(c.get('주요사업'))}",
    ]
    if has_value(c.get("최근활동")):
        parts.append(f"최근활동: {c['최근활동']}")
    if has_value(c.get("제안안")):
        parts.append(f"제안안: {c['제안안']}")
    if urls:
        parts.append("[출처] " + ", ".join(urls))
    return " / ".join(parts)


def contact_row(c, cfg):
    phone, email = val(c.get("전화")), val(c.get("이메일"))
    homepage = str(c.get("홈페이지") or "").strip() or NO_HOMEPAGE
    row = {h: "" for h in CONTACT_HEADERS}
    row.update({
        "번호": c["번호"],
        "기업명": c["기업명"],
        "담당자 연락처(전화번호+이메일)": f"{phone} / {email}",
        "홈페이지 주소": homepage if homepage != UNKNOWN else NO_HOMEPAGE,
        "기업소개 및 제안주제": intro(c),
        "직무명": cfg["team"].get("job_title", ""),
        "담당자명/직급": str(c.get("담당자") or "").strip() or cfg["team"].get("default_contact_title", "담당자"),
    })
    return [row[h] for h in CONTACT_HEADERS]


def verify_row(c):
    return [
        c["번호"], c["기업명"], c.get("시도", ""), c.get("시군구", ""), c.get("업종", ""), val(c.get("임직원수")),
        "②" if tier(c) == 3 else "①", "예" if is_focus(c) else "아니오",
        "예" if str(c.get("재검토")) == "예" else "아니오", c.get("출처사이트", ""),
        "\n".join(u for u in (c.get("출처URL") or []) if u), c.get("비고", ""),
    ]


def summary_rows(summary):
    rows = []
    for k, v in summary.items():
        if isinstance(v, dict):
            for kk, vv in v.items():
                rows.append([k, str(kk), vv if not isinstance(vv, dict) else str(vv)])
        else:
            rows.append([k, "", v])
    return rows


def style(ws, widths, wrap_cols=()):
    for cell in ws[1]:
        cell.font = Font(bold=True)
    ws.freeze_panes = "A2"
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[ws.cell(row=1, column=i).column_letter].width = w
    for col in wrap_cols:
        for cell in ws[ws.cell(row=1, column=col).column_letter]:
            cell.alignment = Alignment(wrap_text=True, vertical="top")


def main(argv=None):
    ap = argparse.ArgumentParser()
    ap.add_argument("--config", required=True)
    ap.add_argument("--run", default=Path("run"), type=Path)
    ap.add_argument("--progress", type=Path)
    ap.add_argument("--existing", type=Path)
    ap.add_argument("--start-number", type=int, help="지정하면 진행기록·명단 대신 이 번호부터")
    ap.add_argument("--out-dir", default=Path("outputs"), type=Path)
    args = ap.parse_args(argv)

    cfg = load_config(args.config)
    run = args.run
    selected = load_json(run / "selected.json")
    if not selected:
        print("선정된 기업이 없어 엑셀을 만들지 않습니다.")
        return 1
    start = args.start_number or next_number(args.progress, args.existing)
    for i, c in enumerate(selected):
        c["번호"] = start + i
    end = start + len(selected) - 1

    _, rejected, _ = load_candidates(run)
    dropped = load_json(run / "dropped.json") if (run / "dropped.json").exists() else []
    dups = load_json(run / "duplicates.json") if (run / "duplicates.json").exists() else []
    summary = load_json(run / "summary.json") if (run / "summary.json").exists() else {}

    wb = Workbook()
    ws = wb.active
    ws.title = SHEET_NAMES[0]
    ws.append(CONTACT_HEADERS)
    for c in selected:
        ws.append(contact_row(c, cfg))
    for row in ws.iter_rows(min_row=2, min_col=9, max_col=9):
        for cell in row:
            cell.number_format = "@"
    style(ws, [6, 8, 12, 24, 10, 10, 12, 10, 34, 30, 10, 10, 70, 16, 16, 10], wrap_cols=[13])

    ws2 = wb.create_sheet(SHEET_NAMES[1])
    ws2.append(VERIFY_HEADERS)
    for c in selected:
        ws2.append(verify_row(c))
    style(ws2, [6, 24, 8, 12, 16, 10, 10, 10, 10, 22, 50, 50], wrap_cols=[11, 12])

    ws3 = wb.create_sheet(SHEET_NAMES[2])
    ws3.append(REJECT_HEADERS)
    for r in rejected:
        ws3.append([r.get("기업명", ""), r.get("사유", ""), r.get("출처", "")])
    for d in dropped:
        ws3.append([d.get("기업명", ""), "원문 재확인 불일치: " + d.get("사유", ""), d.get("출처", "")])
    for d in dups:
        ws3.append([d.get("기업명", ""), d.get("_중복", "중복"), d.get("출처사이트", "")])
    style(ws3, [28, 80, 40], wrap_cols=[2])

    ws4 = wb.create_sheet(SHEET_NAMES[3])
    ws4.append(["항목", "구분", "값"])
    for row in summary_rows(summary):
        ws4.append(row)
    style(ws4, [16, 30, 14])

    stamp = datetime.now(KST).strftime("%Y%m%d_%H%M")
    args.out_dir.mkdir(parents=True, exist_ok=True)
    out = args.out_dir / f"기업컨택_리스트_추가분_{stamp}({start}-{end}).xlsx"
    wb.save(out)

    problems = validate(str(out), cfg.get("run", {}).get("max_per_run", 50))
    if problems:
        print("엑셀 검증 실패:")
        for p in problems:
            print(" -", p)
        return 1
    save_json(run / "numbered.json", selected)
    save_json(run / "xlsx_name.json", {"file": out.name, "start": start, "end": end})
    print(f"엑셀 작성·검증 통과: {out} ({start}~{end}, {len(selected)}행)")
    return 0


if __name__ == "__main__":
    sys.exit(main())
